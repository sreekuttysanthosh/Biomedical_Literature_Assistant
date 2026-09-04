import os
import re
from io import BytesIO
from xml.etree import ElementTree as ET

import streamlit as st
from dotenv import load_dotenv
from google import genai
from Bio import Entrez

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)


# ============================================================
# CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="AI Biomedical Literature Assistant",
    page_icon="🧬",
    layout="wide",
    initial_sidebar_state="expanded",
)

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY:
    st.error(
        "Gemini API key not found. "
        "Please check your .env file."
    )
    st.stop()

client = genai.Client(api_key=GEMINI_API_KEY)

# Required by NCBI Entrez
Entrez.email = "iamkukuluz30@gmail.com"


# ============================================================
# SESSION STATE
# ============================================================

if "papers" not in st.session_state:
    st.session_state.papers = []

if "analysis" not in st.session_state:
    st.session_state.analysis = ""

if "search_topic" not in st.session_state:
    st.session_state.search_topic = ""


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    .stApp {
        background-color: #f6f8fc;
    }

    .block-container {
        max-width: 1200px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    .hero {
        background: linear-gradient(
            135deg,
            #667eea,
            #764ba2
        );
        padding: 40px;
        border-radius: 22px;
        margin-bottom: 30px;
        color: white;
        box-shadow: 0 10px 30px rgba(70, 60, 150, 0.15);
    }

    .hero-title {
        font-size: 38px;
        font-weight: 750;
        margin-bottom: 10px;
        color: white;
    }

    .hero-text {
        font-size: 17px;
        line-height: 1.6;
        color: white;
    }

    .feature-card {
        background: white;
        padding: 25px;
        border-radius: 18px;
        border: 1px solid #e5e7eb;
        min-height: 190px;
        box-shadow: 0 5px 20px rgba(0, 0, 0, 0.04);
    }

    .feature-icon {
        font-size: 32px;
        margin-bottom: 8px;
    }

    .feature-title {
        font-size: 20px;
        font-weight: 700;
        color: #202938;
        margin-bottom: 10px;
    }

    .feature-text {
        font-size: 14px;
        line-height: 1.6;
        color: #667085;
    }

    .paper-card {
        background: white;
        padding: 22px;
        border-radius: 16px;
        border: 1px solid #e5e7eb;
        margin-bottom: 16px;
        box-shadow: 0 4px 15px rgba(0, 0, 0, 0.03);
    }

    .paper-title {
        font-size: 18px;
        font-weight: 700;
        color: #243b64;
        margin-bottom: 10px;
    }

    .paper-meta {
        font-size: 14px;
        color: #667085;
        line-height: 1.7;
    }

    .paper-abstract {
        font-size: 14px;
        line-height: 1.6;
        color: #344054;
        margin-top: 12px;
    }

    .insight-box {
        background: white;
        padding: 25px;
        border-radius: 18px;
        border-left: 5px solid #667eea;
        box-shadow: 0 5px 20px rgba(0, 0, 0, 0.04);
        margin-top: 20px;
    }

    .footer {
        text-align: center;
        color: #667085;
        padding: 30px;
        margin-top: 40px;
    }

    section[data-testid="stSidebar"] {
        background-color: #f1f4f9;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.title("🧬 Biomedical AI")

    st.divider()

    st.subheader("🔎 What this tool does")

    st.markdown("### 1. Search Literature")
    st.write(
        "Search real biomedical research "
        "papers from PubMed."
    )

    st.markdown("### 2. Retrieve Evidence")
    st.write(
        "Collect titles, authors, journals, "
        "abstracts and PMIDs."
    )

    st.markdown("### 3. AI Analysis")
    st.write(
        "Gemini analyzes the retrieved "
        "biomedical literature."
    )

    st.markdown("### 4. Research Insights")
    st.write(
        "Identify findings, limitations "
        "and potential research gaps."
    )

    st.divider()

    st.caption(
        "Python • BioPython • Streamlit • PubMed • Gemini"
    )


# ============================================================
# HERO
# ============================================================

st.markdown(
    """
<div class="hero">
<div class="hero-title">🧬 AI Biomedical Literature Assistant</div>
<div class="hero-text">
Turn biomedical research into clear, actionable insights.
Search PubMed, summarize scientific evidence, compare findings, and uncover potential research gaps with AI
</div>
</div>
""",
    unsafe_allow_html=True,
)

# ============================================================
# FEATURE CARDS
# ============================================================

col1, col2, col3 = st.columns(3)

with col1:
    st.markdown(
        """
<div class="feature-card">
<div class="feature-icon">🔎</div>
<div class="feature-title">Search Literature</div>
<div class="feature-text">
Search real biomedical publications from PubMed
using a research topic.
</div>
</div>
""",
        unsafe_allow_html=True,
    )

with col2:
    st.markdown(
        """
<div class="feature-card">
<div class="feature-icon">🤖</div>
<div class="feature-title">AI Analysis</div>
<div class="feature-text">
Use Generative AI to summarize and analyze
retrieved biomedical literature.
</div>
</div>
""",
        unsafe_allow_html=True,
    )

with col3:
    st.markdown(
        """
<div class="feature-card">
<div class="feature-icon">💡</div>
<div class="feature-title">Research Insights</div>
<div class="feature-text">
Identify findings, limitations and potential
research gaps.
</div>
</div>
""",
        unsafe_allow_html=True,
    )

# ============================================================
# PUBMED SEARCH FUNCTION
# ============================================================

def search_pubmed(search_term, max_results):

    # Search PubMed
    search_handle = Entrez.esearch(
        db="pubmed",
        term=search_term,
        retmax=max_results,
        sort="relevance",
    )

    search_record = Entrez.read(search_handle)
    search_handle.close()

    pmids = search_record.get("IdList", [])

    if not pmids:
        return []

    # Fetch complete records
    fetch_handle = Entrez.efetch(
        db="pubmed",
        id=",".join(pmids),
        retmode="xml",
    )

    xml_data = fetch_handle.read()
    fetch_handle.close()

    root = ET.fromstring(xml_data)

    papers = []

    for article in root.findall(".//PubmedArticle"):

        # ----------------------------------------------------
        # PMID
        # ----------------------------------------------------

        pmid_element = article.find(".//PMID")

        pmid = (
            pmid_element.text
            if pmid_element is not None
            else ""
        )

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title_element = article.find(
            ".//ArticleTitle"
        )

        title = (
            "".join(title_element.itertext())
            if title_element is not None
            else "No title available"
        )

        # ----------------------------------------------------
        # AUTHORS
        # ----------------------------------------------------

        authors = []

        for author in article.findall(
            ".//AuthorList/Author"
        ):

            last_name = author.find("LastName")
            initials = author.find("Initials")

            if last_name is not None:

                name = last_name.text or ""

                if initials is not None:
                    name += " " + (
                        initials.text or ""
                    )

                authors.append(name)

        # ----------------------------------------------------
        # JOURNAL
        # ----------------------------------------------------

        journal_element = article.find(
            ".//Journal/Title"
        )

        journal = (
            journal_element.text
            if journal_element is not None
            else "Unknown journal"
        )

        # ----------------------------------------------------
        # YEAR
        # ----------------------------------------------------

        year = "Unknown year"

        year_element = article.find(
            ".//PubDate/Year"
        )

        if year_element is not None:
            year = year_element.text

        else:

            medline_date = article.find(
                ".//PubDate/MedlineDate"
            )

            if medline_date is not None:
                year = (
                    medline_date.text[:4]
                    if medline_date.text
                    else "Unknown year"
                )

        # ----------------------------------------------------
        # ABSTRACT
        # ----------------------------------------------------

        abstract_parts = []

        for abstract_element in article.findall(
            ".//Abstract/AbstractText"
        ):

            text = "".join(
                abstract_element.itertext()
            )

            label = abstract_element.attrib.get(
                "Label"
            )

            if label:
                text = f"{label}: {text}"

            abstract_parts.append(text)

        abstract = " ".join(
            abstract_parts
        ).strip()

        # ----------------------------------------------------
        # PMID URL
        # ----------------------------------------------------

        pubmed_url = (
            f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"
        )

        papers.append(
            {
                "pmid": pmid,
                "title": title,
                "authors": authors,
                "journal": journal,
                "year": year,
                "abstract": abstract,
                "url": pubmed_url,
            }
        )

    return papers


# ============================================================
# SEARCH INTERFACE
# ============================================================

st.subheader("🔎 Search Biomedical Literature")

st.write(
    "Search for a biomedical topic and retrieve "
    "relevant research papers from PubMed."
)

topic = st.text_input(
    "Biomedical Research Topic",
    value=st.session_state.search_topic,
    placeholder=(
        "Example: antibiotic resistance "
        "in Pseudomonas aeruginosa"
    ),
)

number_of_papers = st.slider(
    "Number of papers",
    min_value=3,
    max_value=10,
    value=5,
)

search_button = st.button(
    "🔎 Search PubMed",
    type="primary",
)


# ============================================================
# SEARCH EXECUTION
# ============================================================

if search_button:

    if not topic.strip():

        st.warning(
            "Please enter a biomedical research topic."
        )

    else:

        st.session_state.search_topic = topic.strip()

        with st.spinner(
            "🔎 Searching PubMed..."
        ):

            try:

                papers = search_pubmed(
                    topic.strip(),
                    number_of_papers,
                )

                st.session_state.papers = papers
                st.session_state.analysis = ""

                if papers:

                    st.success(
                        f"Found {len(papers)} "
                        "biomedical publications."
                    )

                else:

                    st.warning(
                        "No PubMed papers were found "
                        "for this topic."
                    )

            except Exception as error:

                st.error(
                    f"PubMed search failed: {error}"
                )


# ============================================================
# DISPLAY PAPERS
# ============================================================

if st.session_state.papers:

    st.subheader("📚 Retrieved Publications")

    for index, paper in enumerate(
        st.session_state.papers,
        start=1,
    ):

        st.markdown(
            '<div class="paper-card">',
            unsafe_allow_html=True,
        )

        st.markdown(
            f"""
            <div class="paper-title">
                {index}. {paper["title"]}
            </div>
            """,
            unsafe_allow_html=True,
        )

        authors_text = ", ".join(
            paper["authors"]
        )

        st.markdown(
            f"""
            <div class="paper-meta">
                <b>Authors:</b>
                {authors_text if authors_text else "Not available"}
                <br>
                <b>Journal:</b> {paper["journal"]}
                <br>
                <b>Year:</b> {paper["year"]}
                <br>
                <b>PMID:</b> {paper["pmid"]}
            </div>
            """,
            unsafe_allow_html=True,
        )

        if paper["abstract"]:

            st.markdown(
                f"""
                <div class="paper-abstract">
                    <b>Abstract:</b><br>
                    {paper["abstract"]}
                </div>
                """,
                unsafe_allow_html=True,
            )

        else:

            st.info(
                "Abstract not available."
            )

        st.link_button(
            "🔗 View on PubMed",
            paper["url"],
        )

        st.markdown(
            "</div>",
            unsafe_allow_html=True,
        )


# ============================================================
# GEMINI ANALYSIS
# ============================================================

if st.session_state.papers:

    st.subheader(
        "🤖 AI-Powered Literature Analysis"
    )

    st.write(
        "Gemini analyzes the retrieved PubMed evidence "
        "and provides a research-oriented summary."
    )

    analyze_button = st.button(
        "🤖 Analyze Literature with Gemini",
        type="primary",
    )

    if analyze_button:

        evidence_parts = []

        for i, paper in enumerate(
            st.session_state.papers,
            start=1,
        ):

            evidence_parts.append(
                f"""
PAPER {i}

Title:
{paper["title"]}

Authors:
{", ".join(paper["authors"])}

Journal:
{paper["journal"]}

Year:
{paper["year"]}

PMID:
{paper["pmid"]}

Abstract:
{paper["abstract"]}

--------------------------------
"""
            )

        evidence = "\n".join(
            evidence_parts
        )

        analysis_prompt = f"""
You are an AI biomedical literature research assistant.

Analyze ONLY the biomedical evidence supplied below.

RESEARCH TOPIC:
{st.session_state.search_topic}

RETRIEVED PUBMED PAPERS:
{evidence}

Provide the following sections:

## 1. Overall Literature Summary

Give a concise overview of what the retrieved
literature says about the research topic.

## 2. Key Findings

Identify the major findings reported
across the retrieved papers.

## 3. Common Themes

Identify recurring biological, molecular,
clinical or technological themes.

## 4. Research Methods

Briefly describe the major methods or
approaches mentioned in the papers.

## 5. Limitations

Identify limitations that are explicitly
supported by the provided abstracts.

Do NOT invent limitations.

## 6. Potential Research Gaps

Identify possible gaps based on the
retrieved evidence.

Clearly label these as
POTENTIAL research gaps.

## 7. Future Research Directions

Suggest reasonable future research
directions based on the evidence.

Clearly distinguish suggestions
from reported findings.

## 8. Important Note

State that this is an AI-assisted
literature analysis and that the
original publications should be
consulted before making scientific
or clinical decisions.

IMPORTANT RULES:

- Do not fabricate findings.
- Do not invent experimental results.
- Do not claim information that is not
  supported by the supplied literature.
- Clearly distinguish evidence from suggestions.
- Use scientific but understandable language.
"""

        with st.spinner(
            "🤖 Gemini is analyzing the literature..."
        ):

            try:

                response = client.models.generate_content(
                    model="gemini-3.5-flash-lite",
                    contents=analysis_prompt,
                )

                st.session_state.analysis = (
                    response.text
                )

            except Exception as error:

                st.error(
                    f"Gemini analysis failed: {error}"
                )


# ============================================================
# WORD DOCUMENT GENERATION
# ============================================================

def create_word_document():

    document = Document()

    title = document.add_heading(
        "AI Biomedical Literature Analysis",
        level=0,
    )

    document.add_paragraph(
        f"Research Topic: "
        f"{st.session_state.search_topic}"
    )

    document.add_paragraph(
        "Generated using AI Biomedical "
        "Literature Assistant"
    )

    document.add_heading(
        "Retrieved Publications",
        level=1,
    )

    for i, paper in enumerate(
        st.session_state.papers,
        start=1,
    ):

        document.add_heading(
            f"{i}. {paper['title']}",
            level=2,
        )

        document.add_paragraph(
            f"Authors: "
            f"{', '.join(paper['authors'])}"
        )

        document.add_paragraph(
            f"Journal: {paper['journal']}"
        )

        document.add_paragraph(
            f"Year: {paper['year']}"
        )

        document.add_paragraph(
            f"PMID: {paper['pmid']}"
        )

        document.add_paragraph(
            f"PubMed: {paper['url']}"
        )

        if paper["abstract"]:

            document.add_paragraph(
                "Abstract:"
            )

            document.add_paragraph(
                paper["abstract"]
            )

    document.add_page_break()

    document.add_heading(
        "AI Literature Analysis",
        level=1,
    )

    document.add_paragraph(
        st.session_state.analysis
    )

    document.add_paragraph(
        "Note: This document contains "
        "AI-assisted literature analysis. "
        "Consult the original publications "
        "before making scientific or clinical decisions."
    )

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# PDF GENERATION
# ============================================================

def clean_text_for_pdf(text):

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text,
    )

    text = re.sub(
        r"#+\s*",
        "",
        text,
    )

    text = text.replace(
        "•",
        "-",
    )

    return text


def create_pdf():

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )

    styles = getSampleStyleSheet()

    story = []

    story.append(
        Paragraph(
            "AI Biomedical Literature Analysis",
            styles["Title"],
        )
    )

    story.append(
        Spacer(1, 0.2 * inch)
    )

    story.append(
        Paragraph(
            f"Research Topic: "
            f"{st.session_state.search_topic}",
            styles["Normal"],
        )
    )

    story.append(
        Spacer(1, 0.3 * inch)
    )

    story.append(
        Paragraph(
            "Retrieved Publications",
            styles["Heading1"],
        )
    )

    for i, paper in enumerate(
        st.session_state.papers,
        start=1,
    ):

        story.append(
            Paragraph(
                f"{i}. {paper['title']}",
                styles["Heading2"],
            )
        )

        story.append(
            Paragraph(
                f"Authors: "
                f"{', '.join(paper['authors'])}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"Journal: {paper['journal']}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"Year: {paper['year']}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                f"PMID: {paper['pmid']}",
                styles["Normal"],
            )
        )

        if paper["abstract"]:

            story.append(
                Paragraph(
                    "Abstract:",
                    styles["Heading3"],
                )
            )

            abstract_text = clean_text_for_pdf(
                paper["abstract"]
            )

            story.append(
                Paragraph(
                    abstract_text,
                    styles["Normal"],
                )
            )

        story.append(
            Spacer(1, 0.2 * inch)
        )

    story.append(
        Paragraph(
            "AI Literature Analysis",
            styles["Heading1"],
        )
    )

    analysis_text = clean_text_for_pdf(
        st.session_state.analysis
    )

    paragraphs = analysis_text.split("\n")

    for paragraph in paragraphs:

        paragraph = paragraph.strip()

        if paragraph:

            story.append(
                Paragraph(
                    paragraph,
                    styles["Normal"],
                )
            )

            story.append(
                Spacer(1, 0.08 * inch)
            )

    story.append(
        Spacer(1, 0.2 * inch)
    )

    story.append(
        Paragraph(
            "Note: This is an AI-assisted "
            "literature analysis. Consult the "
            "original publications before making "
            "scientific or clinical decisions.",
            styles["Normal"],
        )
    )

    document.build(story)

    output.seek(0)

    return output.getvalue()


# ============================================================
# DOWNLOAD SECTION
# ============================================================

if st.session_state.analysis:

    st.subheader("💡 Research Insights")

    st.markdown(
        '<div class="insight-box">',
        unsafe_allow_html=True,
    )

    st.markdown(
        st.session_state.analysis
    )

    st.markdown(
        "</div>",
        unsafe_allow_html=True,
    )

    st.subheader(
        "📥 Download Your Analysis"
    )

    col1, col2 = st.columns(2)

    with col1:

        word_file = create_word_document()

        st.download_button(
            label="📄 Download Word Document",
            data=word_file,
            file_name=(
                "biomedical_literature_analysis.docx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
        )

    with col2:

        pdf_file = create_pdf()

        st.download_button(
            label="📕 Download PDF",
            data=pdf_file,
            file_name=(
                "biomedical_literature_analysis.pdf"
            ),
            mime="application/pdf",
        )


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.markdown(
    """
<div class="footer">
<b>🧬 AI Biomedical Literature Assistant</b>
<br><br>
Built using Python • BioPython • Streamlit • PubMed • Google Gemini
<br><br>
AI-assisted biomedical literature exploration and analysis.
</div>
""",
    unsafe_allow_html=True,
)